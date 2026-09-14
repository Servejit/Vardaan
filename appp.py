# ============================================================
# 6THSENSE (6S-FO200) VARDAAN
# STREAMLIT + SUPABASE
#
# MASTER EXCEL
#       ↓
# VALUE-ONLY OUTPUT
#       ↓
# 7 CONDITIONS + BLUE/GREEN COLORS
#       ↓
# 6thsense(6S-FO200)Vardaan.xlsx
#
# SUPABASE AUTH:
#   Admin = UUID exists in admin_users
#   User  = authenticated user but not in admin_users
#
# ADMIN:
#   Upload Master
#   Generate Output
#   Clear Master
#
# NORMAL USER:
#   Download generated output only
# ============================================================

import streamlit as st
import pandas as pd
import numpy as np
import openpyxl

import os
import re
import base64
import io
import json
import hashlib
import requests

from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

from supabase import create_client


# ============================================================
#                    STREAMLIT CONFIG
# ============================================================

st.set_page_config(
    page_title="6thSense (6S-FO200) Vardaan",
    page_icon="📊",
    layout="wide"
)


# ============================================================
#                    SUPABASE SETTINGS
# ============================================================

SUPABASE_URL = "https://uzekepedksagskhcpwhk.supabase.co"

# IMPORTANT:
# Put your existing sb_publishable key here,
# OR preferably put it in Streamlit Secrets.
#
# .streamlit/secrets.toml:
#
# SUPABASE_URL = "https://uzekepedksagskhcpwhk.supabase.co"
# SUPABASE_ANON_KEY = "your_existing_publishable_key"
#
# The code below first checks Secrets.
# ============================================================

try:
    SUPABASE_URL = st.secrets.get(
        "SUPABASE_URL",
        SUPABASE_URL
    )

    SUPABASE_ANON_KEY = st.secrets.get(
        "SUPABASE_ANON_KEY",
        ""
    )

except Exception:
    SUPABASE_ANON_KEY = ""


if not SUPABASE_ANON_KEY:

    st.error(
        "SUPABASE_ANON_KEY is missing from Streamlit Secrets."
    )

    st.stop()


supabase = create_client(
    SUPABASE_URL,
    SUPABASE_ANON_KEY
)
def restore_supabase_session():

    try:
        session = st.session_state.get(
            "auth_session"
        )

        if session is None:
            return False

        access_token = session.access_token
        refresh_token = session.refresh_token

        if not access_token:
            return False

        # Restore Supabase Auth session
        supabase.auth.set_session(
            access_token,
            refresh_token
        )

        # Explicitly attach JWT to database requests
        supabase.postgrest.auth(
            access_token
        )

        return True

    except Exception:
        return False

# ============================================================
#                    CONSTANTS
# ============================================================

OUTPUT_NAME = "6thsense(6S-FO200)Vardaan.xlsx"

MASTER_RECORD_NAME = "MASTER_FILE"

OUTPUT_RECORD_NAME = OUTPUT_NAME


DEFAULT_BLUE = "ADD8E6"
DEFAULT_GREEN = "90EE90"


# ============================================================
#                    SESSION STATE
# ============================================================

if "user" not in st.session_state:
    st.session_state.user = None

if "is_admin" not in st.session_state:
    st.session_state.is_admin = False

if "message" not in st.session_state:
    st.session_state.message = ""

if "error" not in st.session_state:
    st.session_state.error = ""


# ============================================================
#                    HELPER FUNCTIONS
# ============================================================

def normalize_header(value):

    if value is None:
        return ""

    return re.sub(
        r"[^a-z0-9]+",
        "",
        str(value).strip().lower()
    )


def numeric_value(value):

    if value is None:
        return None

    if isinstance(value, bool):
        return None

    if isinstance(value, (int, float, np.integer, np.floating)):

        if pd.isna(value):
            return None

        return float(value)

    text = str(value).strip()

    if not text:
        return None

    # Remove commas
    text = text.replace(",", "")

    # Remove percentage sign
    text = text.replace("%", "")

    # Remove common currency symbols
    text = text.replace("₹", "")
    text = text.replace("$", "")

    try:
        return float(text)

    except Exception:

        # Try to find a number in text
        match = re.search(
            r"[-+]?\d*\.?\d+",
            text
        )

        if match:

            try:
                return float(match.group())

            except Exception:
                return None

        return None


def first_number(value):

    if value is None:
        return None

    text = str(value)

    match = re.search(
        r"[-+]?\d*\.?\d+",
        text
    )

    if not match:
        return None

    try:
        return float(match.group())

    except Exception:
        return None


def number_inside_parentheses(value):

    if value is None:
        return None

    text = str(value)

    match = re.search(
        r"\(\s*([-+]?\d*\.?\d+)\s*\)",
        text
    )

    if not match:
        return None

    try:
        return float(match.group(1))

    except Exception:
        return None


def find_heading(ws, heading):

    target = normalize_header(heading)

    for cell in ws[1]:

        if normalize_header(cell.value) == target:
            return cell.column

    return None


def find_heading_contains(ws, text):

    target = normalize_header(text)

    for cell in ws[1]:

        if target in normalize_header(cell.value):
            return cell.column

    return None


# ============================================================
#                    DATE DETECTION
# ============================================================

MONTHS = {
    "jan": 1,
    "january": 1,

    "feb": 2,
    "february": 2,

    "mar": 3,
    "march": 3,

    "apr": 4,
    "april": 4,

    "may": 5,

    "jun": 6,
    "june": 6,

    "jul": 7,
    "july": 7,

    "aug": 8,
    "august": 8,

    "sep": 9,
    "sept": 9,
    "september": 9,

    "oct": 10,
    "october": 10,

    "nov": 11,
    "november": 11,

    "dec": 12,
    "december": 12
}


def extract_date_from_heading(value):

    if value is None:
        return None

    # Actual datetime
    if isinstance(value, datetime):
        return value

    # pandas Timestamp
    if isinstance(value, pd.Timestamp):

        if pd.isna(value):
            return None

        return value.to_pydatetime()

    text = str(value).strip()

    if not text:
        return None

    # --------------------------------------------------------
    # 10 Sep 2026
    # 10 September 2026
    # --------------------------------------------------------

    match = re.search(
        r"\b(\d{1,2})\s+"
        r"([A-Za-z]+)"
        r"(?:\s+(\d{2,4}))?\b",
        text
    )

    if match:

        day = int(match.group(1))

        month_text = match.group(2).lower()

        month = MONTHS.get(
            month_text
        )

        year_text = match.group(3)

        if month:

            if year_text:

                year = int(year_text)

                if year < 100:
                    year += 2000

            else:

                year = datetime.now().year

            try:

                return datetime(
                    year,
                    month,
                    day
                )

            except Exception:
                pass

    # --------------------------------------------------------
    # 10-09-2026
    # 10/09/2026
    # 10.09.2026
    # --------------------------------------------------------

    match = re.search(
        r"\b(\d{1,2})[-/.]"
        r"(\d{1,2})[-/.]"
        r"(\d{2,4})\b",
        text
    )

    if match:

        day = int(match.group(1))
        month = int(match.group(2))
        year = int(match.group(3))

        if year < 100:
            year += 2000

        try:

            return datetime(
                year,
                month,
                day
            )

        except Exception:
            pass

    return None


def find_recent_p2l_column(ws):

    candidates = []

    for cell in ws[1]:

        value = cell.value

        normalized = normalize_header(value)

        if (
            "p2l" not in normalized
            and "o2l" not in normalized
        ):
            continue

        date_value = extract_date_from_heading(
            value
        )

        if date_value is not None:

            candidates.append(
                (
                    date_value,
                    cell.column
                )
            )

    if not candidates:
        return None

    candidates.sort(
        key=lambda x: x[0],
        reverse=True
    )

    return candidates[0][1]


# ============================================================
#                    COLOR FUNCTIONS
# ============================================================

def get_fill(hex_color):

    return PatternFill(
        fill_type="solid",
        fgColor=str(
            hex_color
        ).replace(
            "#",
            ""
        ).upper()
    )


def set_cell_blue(
    cell,
    blue_fill
):

    cell.fill = blue_fill


def set_row_green(
    ws,
    row_number,
    green_fill
):

    for col in range(
        1,
        ws.max_column + 1
    ):

        ws.cell(
            row=row_number,
            column=col
        ).fill = green_fill


# ============================================================
#                    SUPABASE HELPERS
# ============================================================

def get_current_user():

    try:

        response = supabase.auth.get_user()

        if response and response.user:
            return response.user

    except Exception:
        pass

    return None


def is_admin_user(user_id):

    if not user_id:
        return False

    try:

        response = (
            supabase
            .table("admin_users")
            .select("user_id")
            .eq("user_id", user_id)
            .limit(1)
            .execute()
        )

        return bool(
            response.data
        )

    except Exception as e:

        st.error(
            f"Unable to check admin status: {e}"
        )

        return False


def save_file_to_supabase(
    file_name,
    file_bytes,
    uploaded_by=None
):

    encoded = base64.b64encode(
        file_bytes
    ).decode("utf-8")

    # Check whether record already exists
    existing = (
        supabase
        .table("master_file")
        .select("id")
        .eq(
            "file_name",
            file_name
        )
        .limit(1)
        .execute()
    )

    row = {
        "file_name": file_name,
        "file_data": encoded,
        "uploaded_by": uploaded_by,
        "uploaded_at": datetime.utcnow().isoformat(),
        "updated_at": datetime.utcnow().isoformat()
    }

    if existing.data:

        record_id = existing.data[0]["id"]

        response = (
            supabase
            .table("master_file")
            .update(row)
            .eq(
                "id",
                record_id
            )
            .execute()
        )

    else:

        response = (
            supabase
            .table("master_file")
            .insert(row)
            .execute()
        )

    return response


def get_file_from_supabase(
    file_name
):

    response = (
        supabase
        .table("master_file")
        .select(
            "id,file_name,file_data,uploaded_at"
        )
        .eq(
            "file_name",
            file_name
        )
        .order(
            "updated_at",
            desc=True
        )
        .limit(1)
        .execute()
    )

    if not response.data:
        return None

    record = response.data[0]

    try:

        return base64.b64decode(
            record["file_data"]
        )

    except Exception:

        return None


def delete_file_from_supabase(
    file_name
):

    response = (
        supabase
        .table("master_file")
        .delete()
        .eq(
            "file_name",
            file_name
        )
        .execute()
    )

    return response


# ============================================================
#                    MASTER UPLOAD
# ============================================================

def upload_master_to_supabase(
    uploaded_file,
    user_id
):

    if uploaded_file is None:
        return False, "Please select a Master Excel file."

    file_name = uploaded_file.name

    allowed_extensions = (
        ".xlsx",
        ".xlsm",
        ".xltx",
        ".xltm"
    )

    if not file_name.lower().endswith(
        allowed_extensions
    ):

        return False, (
            "Only Excel files are allowed."
        )

    file_bytes = uploaded_file.getvalue()

    if not file_bytes:
        return False, "The uploaded file is empty."

    try:

        # Validate workbook before saving
        test_wb = load_workbook(
            io.BytesIO(file_bytes),
            data_only=True
        )

        if "summary" not in test_wb.sheetnames:

            test_wb.close()

            return False, (
                "The Master Excel must contain "
                "a sheet named 'summary'."
            )

        test_wb.close()

    except Exception as e:

        return False, (
            f"Unable to read Master Excel: {e}"
        )

    try:

        save_file_to_supabase(
            MASTER_RECORD_NAME,
            file_bytes,
            user_id
        )

        return True, (
            "Master Excel uploaded successfully."
        )

    except Exception as e:

        return False, (
            f"Supabase upload failed: {e}"
        )


# ============================================================
#                    CREATE OUTPUT
# ============================================================

def create_output(
    master_bytes
):

    if not master_bytes:

        raise Exception(
            "Master Excel is not available."
        )

    # ========================================================
    # LOAD SOURCE WORKBOOK
    # ========================================================

    source_wb = load_workbook(
        io.BytesIO(master_bytes),
        data_only=True
    )

    # Original code requires "summary"
    if "summary" not in source_wb.sheetnames:

        source_wb.close()

        raise Exception(
            "Master Excel does not contain "
            "a sheet named 'summary'."
        )

    source_ws = source_wb["summary"]

    # ========================================================
    # CREATE NEW WORKBOOK
    # ========================================================

    output_wb = Workbook()

    output_ws = output_wb.active

    output_ws.title = "summary"

    # ========================================================
    # COPY ONLY CELL VALUES
    #
    # This preserves the original Colab behavior:
    # formulas are not copied.
    # Cached values are copied.
    # ========================================================

    for row in source_ws.iter_rows():

        for source_cell in row:

            output_ws.cell(
                row=source_cell.row,
                column=source_cell.column
            ).value = source_cell.value

    # ========================================================
    # FREEZE / FILTER
    # ========================================================

    output_ws.freeze_panes = "A2"

    if (
        output_ws.max_row >= 1
        and output_ws.max_column >= 1
    ):

        output_ws.auto_filter.ref = (
            output_ws.dimensions
        )

    # ========================================================
    # HEADER FORMATTING
    # ========================================================

    for cell in output_ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )

    # ========================================================
    # FIND REQUIRED COLUMNS
    # ========================================================

    col_sum_i = find_heading(
        output_ws,
        "Sum I"
    )

    col_16gt = find_heading(
        output_ws,
        "16> C-B / Avg.4"
    )

    col_16lt = find_heading(
        output_ws,
        "16< D-B / Avg.4"
    )

    col_o2h = find_heading(
        output_ws,
        "Sum O2H.10"
    )

    col_o2l = find_heading(
        output_ws,
        "Sum O2L.10"
    )

    col_recent_p2l = find_recent_p2l_column(
        output_ws
    )

    # ========================================================
    # COLOR SETTINGS
    # ========================================================

    light_blue = DEFAULT_BLUE
    light_green = DEFAULT_GREEN

    blue_fill = get_fill(
        light_blue
    )

    green_fill = get_fill(
        light_green
    )

    # ========================================================
    # CONDITION COUNTER
    # ========================================================

    condition_count = {}

    for row in range(
        2,
        output_ws.max_row + 1
    ):

        condition_count[row] = 0

    # ========================================================
    # 1. Sum I < -4.00
    # ========================================================

    if col_sum_i:

        for row in range(
            2,
            output_ws.max_row + 1
        ):

            value = numeric_value(
                output_ws.cell(
                    row=row,
                    column=col_sum_i
                ).value
            )

            if (
                value is not None
                and value < -4.00
            ):

                output_ws.cell(
                    row=row,
                    column=col_sum_i
                ).fill = blue_fill

                output_ws.cell(
                    row=row,
                    column=1
                ).fill = blue_fill

                condition_count[row] += 1

    # ========================================================
    # 2. 16> C-B / Avg.4
    #
    # Example:
    # 16>0.25, Avg.4(0.44)
    #
    # Parentheses value < 0.50
    # ========================================================

    if col_16gt:

        for row in range(
            2,
            output_ws.max_row + 1
        ):

            cell = output_ws.cell(
                row=row,
                column=col_16gt
            )

            avg_value = number_inside_parentheses(
                cell.value
            )

            if (
                avg_value is not None
                and avg_value < 0.50
            ):

                cell.fill = blue_fill

                output_ws.cell(
                    row=row,
                    column=1
                ).fill = blue_fill

                condition_count[row] += 1

    # ========================================================
    # 3. 16< D-B / Avg.4
    #
    # Parentheses value < first value
    # AND parentheses value < -1.00
    # ========================================================

    if col_16lt:

        for row in range(
            2,
            output_ws.max_row + 1
        ):

            cell = output_ws.cell(
                row=row,
                column=col_16lt
            )

            first_val = first_number(
                cell.value
            )

            avg_val = number_inside_parentheses(
                cell.value
            )

            if (
                first_val is not None
                and avg_val is not None
                and avg_val < first_val
                and avg_val < -1.00
            ):

                cell.fill = blue_fill

                output_ws.cell(
                    row=row,
                    column=1
                ).fill = blue_fill

                condition_count[row] += 1

    # ========================================================
    # 4. Sum O2H.10
    #
    # Calculate average.
    # Color values BELOW average.
    # ========================================================

    o2h_values = []

    if col_o2h:

        for row in range(
            2,
            output_ws.max_row + 1
        ):

            value = numeric_value(
                output_ws.cell(
                    row=row,
                    column=col_o2h
                ).value
            )

            if value is not None:
                o2h_values.append(value)

    o2h_average = (
        sum(o2h_values) / len(o2h_values)
        if o2h_values
        else None
    )

    if (
        col_o2h
        and o2h_average is not None
    ):

        for row in range(
            2,
            output_ws.max_row + 1
        ):

            value = numeric_value(
                output_ws.cell(
                    row=row,
                    column=col_o2h
                ).value
            )

            if (
                value is not None
                and value < o2h_average
            ):

                output_ws.cell(
                    row=row,
                    column=col_o2h
                ).fill = blue_fill

                output_ws.cell(
                    row=row,
                    column=1
                ).fill = blue_fill

                condition_count[row] += 1

    # ========================================================
    # 5. Sum O2L.10
    #
    # Calculate average.
    # Color values BELOW average.
    # ========================================================

    o2l_values = []

    if col_o2l:

        for row in range(
            2,
            output_ws.max_row + 1
        ):

            value = numeric_value(
                output_ws.cell(
                    row=row,
                    column=col_o2l
                ).value
            )

            if value is not None:
                o2l_values.append(value)

    o2l_average = (
        sum(o2l_values) / len(o2l_values)
        if o2l_values
        else None
    )

    if (
        col_o2l
        and o2l_average is not None
    ):

        for row in range(
            2,
            output_ws.max_row + 1
        ):

            value = numeric_value(
                output_ws.cell(
                    row=row,
                    column=col_o2l
                ).value
            )

            if (
                value is not None
                and value < o2l_average
            ):

                output_ws.cell(
                    row=row,
                    column=col_o2l
                ).fill = blue_fill

                output_ws.cell(
                    row=row,
                    column=1
                ).fill = blue_fill

                condition_count[row] += 1

    # ========================================================
    # 6. MOST RECENT DATE P2L/O2L
    #
    # Condition:
    # value < -1.00
    # ========================================================

    recent_heading = None

    if col_recent_p2l:

        recent_heading = output_ws.cell(
            row=1,
            column=col_recent_p2l
        ).value

        for row in range(
            2,
            output_ws.max_row + 1
        ):

            value = numeric_value(
                output_ws.cell(
                    row=row,
                    column=col_recent_p2l
                ).value
            )

            if (
                value is not None
                and value < -1.00
            ):

                output_ws.cell(
                    row=row,
                    column=col_recent_p2l
                ).fill = blue_fill

                output_ws.cell(
                    row=row,
                    column=1
                ).fill = blue_fill

                condition_count[row] += 1

    # ========================================================
    # 7. FIND ROW(S) WITH MOST BLUE CONDITIONS
    #
    # Entire strongest row becomes LIGHT GREEN.
    #
    # Ties are all colored green.
    # ========================================================

    valid_counts = [
        count
        for count in condition_count.values()
        if count > 0
    ]

    if valid_counts:

        maximum_count = max(
            valid_counts
        )

        strongest_rows = [
            row
            for row, count
            in condition_count.items()
            if count == maximum_count
        ]

        for row in strongest_rows:

            set_row_green(
                output_ws,
                row,
                green_fill
            )

    else:

        maximum_count = 0
        strongest_rows = []

    # ========================================================
    # COLUMN WIDTH
    # ========================================================

    for col in range(
        1,
        output_ws.max_column + 1
    ):

        max_length = 0

        for row in range(
            1,
            min(
                output_ws.max_row,
                1000
            ) + 1
        ):

            value = output_ws.cell(
                row=row,
                column=col
            ).value

            if value is not None:

                length = len(
                    str(value)
                )

                if length > max_length:
                    max_length = length

        output_ws.column_dimensions[
            get_column_letter(col)
        ].width = min(
            max(
                max_length + 2,
                10
            ),
            35
        )

    # ========================================================
    # SAVE TO MEMORY
    # ========================================================

    output_buffer = io.BytesIO()

    output_wb.save(
        output_buffer
    )

    output_wb.close()
    source_wb.close()

    output_buffer.seek(0)

    return (
        output_buffer.getvalue(),
        {
            "o2h_average": o2h_average,
            "o2l_average": o2l_average,
            "recent_heading": recent_heading,
            "maximum_count": maximum_count,
            "strongest_rows": strongest_rows
        }
    )


# ============================================================
#                    GENERATE OUTPUT
# ============================================================

def generate_output():

    master_bytes = get_file_from_supabase(
        MASTER_RECORD_NAME
    )

    if not master_bytes:

        return False, (
            "No Master Excel is currently stored."
        ), None, None

    try:

        output_bytes, result = create_output(
            master_bytes
        )

    except Exception as e:

        return False, (
            f"Excel conversion failed: {e}"
        ), None, None

    try:

        save_file_to_supabase(
            OUTPUT_RECORD_NAME,
            output_bytes,
            st.session_state.user.id
        )

    except Exception as e:

        return False, (
            f"Generated Excel could not be saved "
            f"to Supabase: {e}"
        ), None, None

    return True, (
        "Excel generated successfully."
    ), output_bytes, result


# ============================================================
#                    CLEAR MASTER
# ============================================================

def clear_master():

    errors = []

    try:

        delete_file_from_supabase(
            MASTER_RECORD_NAME
        )

    except Exception as e:

        errors.append(
            f"Master: {e}"
        )

    try:

        delete_file_from_supabase(
            OUTPUT_RECORD_NAME
        )

    except Exception as e:

        errors.append(
            f"Output: {e}"
        )

    if errors:

        return False, (
            "Clear completed with errors: "
            + " | ".join(errors)
        )

    return True, (
        "Master and generated output cleared."
    )


# ============================================================
#                    LOGIN
# ============================================================

def login_user(email, password):

    try:

        response = (
            supabase
            .auth
            .sign_in_with_password(
                {
                    "email": email,
                    "password": password
                }
            )
        )

        if not response.user:
            return False, "Invalid email or password."

        if not response.session:
            return False, "Login session could not be created."

        # Store user
        st.session_state.user = response.user

        # Store complete Auth session
        st.session_state.auth_session = response.session

        # Attach access token to Supabase database requests
        supabase.auth.set_session(
            response.session.access_token,
            response.session.refresh_token
        )

        supabase.postgrest.auth(
            response.session.access_token
        )

        # Check admin
        st.session_state.is_admin = (
            is_admin_user(
                response.user.id
            )
        )

        return True, "Login successful."

    except Exception as e:

        return False, str(e)


# ============================================================
#                    LOGIN SCREEN
# ============================================================

if st.session_state.user is None:

    st.title(
        "6thSense (6S-FO200) Vardaan"
    )

    st.subheader(
        "Login"
    )

    with st.form(
        "login_form"
    ):

        email = st.text_input(
            "Email"
        )

        password = st.text_input(
            "Password",
            type="password"
        )

        login_button = st.form_submit_button(
            "Login",
            use_container_width=True
        )

        if login_button:

            if not email or not password:

                st.error(
                    "Enter email and password."
                )

            else:

                ok, msg = login_user(
                    email,
                    password
                )

                if ok:

                    st.success(msg)

                    st.rerun()

                else:

                    st.error(msg)

    st.stop()


# ============================================================
#                    LOGGED-IN USER
# ============================================================

user = st.session_state.user

st.title(
    "6thSense (6S-FO200) Vardaan"
)

st.caption(
    f"Logged in as: {user.email}"
)

if st.session_state.is_admin:

    st.success(
        "ADMIN ACCESS"
    )

else:

    st.info(
        "USER ACCESS — Read-only"
    )

# ============================================================
#                    LOGOUT FUNCTION
# ============================================================

def logout_user():

    try:
        supabase.auth.sign_out()
    except Exception:
        pass

    st.session_state.user = None
    st.session_state.auth_session = None
    st.session_state.is_admin = False

    st.rerun()


# ============================================================
#                         SIDEBAR
# ============================================================

with st.sidebar:

    st.write(
        f"**User:** {user.email}"
    )

    if st.session_state.is_admin:

        st.write(
            "**Role:** Admin"
        )

    else:

        st.write(
            "**Role:** User"
        )

    if st.button(
        "Logout",
        use_container_width=True
    ):

        logout_user()


# ============================================================
#                    ADMIN PANEL
# ============================================================

if st.session_state.is_admin:

    st.header(
        "Admin Control Panel"
    )

    # --------------------------------------------------------
    # UPLOAD MASTER
    # --------------------------------------------------------

    st.subheader(
        "1. Upload / Replace Master Excel"
    )

    uploaded_file = st.file_uploader(
        "Select Master Excel",
        type=[
            "xlsx",
            "xlsm",
            "xltx",
            "xltm"
        ],
        key="master_uploader"
    )

    if st.button(
        "Upload Master",
        type="primary",
        use_container_width=True
    ):

        if uploaded_file is None:

            st.warning(
                "Please select a Master Excel first."
            )

        else:

            with st.spinner(
                "Uploading Master Excel..."
            ):

                ok, msg = upload_master_to_supabase(
                    uploaded_file,
                    user.id
                )

            if ok:

                st.success(msg)

            else:

                st.error(msg)

    st.divider()

    # --------------------------------------------------------
    # GENERATE
    # --------------------------------------------------------

    st.subheader(
        "2. Generate 6thsense Excel"
    )

    st.write(
        f"Output filename: `{OUTPUT_NAME}`"
    )

    if st.button(
        "Generate Excel",
        type="primary",
        use_container_width=True
    ):

        with st.spinner(
            "Generating value-only Excel and applying conditions..."
        ):

            ok, msg, output_bytes, result = (
                generate_output()
            )

        if ok:

            st.success(msg)

            if result:

                col1, col2 = st.columns(2)

                with col1:

                    if result["o2h_average"] is not None:

                        st.metric(
                            "Sum O2H.10 Average",
                            f"{result['o2h_average']:.4f}"
                        )

                    if result["o2l_average"] is not None:

                        st.metric(
                            "Sum O2L.10 Average",
                            f"{result['o2l_average']:.4f}"
                        )

                with col2:

                    st.metric(
                        "Highest Condition Count",
                        result["maximum_count"]
                    )

                    if result["recent_heading"]:

                        st.write(
                            "**Recent P2L/O2L:**",
                            result["recent_heading"]
                        )

            if output_bytes:

                st.download_button(
                    label=(
                        "Download "
                        "6thsense(6S-FO200)Vardaan.xlsx"
                    ),
                    data=output_bytes,
                    file_name=OUTPUT_NAME,
                    mime=(
                        "application/vnd.openxmlformats-"
                        "officedocument.spreadsheetml.sheet"
                    ),
                    use_container_width=True
                )

        else:

            st.error(msg)

    st.divider()

    # --------------------------------------------------------
    # CLEAR
    # --------------------------------------------------------

    st.subheader(
        "3. Clear Stored Files"
    )

    st.warning(
        "This removes both the stored Master and "
        "generated output from Supabase."
    )

    if st.button(
        "Clear Master + Generated Output",
        type="secondary",
        use_container_width=True
    ):

        ok, msg = clear_master()

        if ok:
            st.success(msg)
        else:
            st.error(msg)


# ============================================================
#                    NORMAL USER PANEL
# ============================================================

else:

    st.header(
        "Generated Excel"
    )

    st.write(
        "You have read-only access. "
        "Only an administrator can upload or replace the Master Excel."
    )

    output_bytes = get_file_from_supabase(
        OUTPUT_RECORD_NAME
    )

    if output_bytes:

        st.success(
            "Generated Excel is available."
        )

        st.download_button(
            label=(
                "Download "
                "6thsense(6S-FO200)Vardaan.xlsx"
            ),
            data=output_bytes,
            file_name=OUTPUT_NAME,
            mime=(
                "application/vnd.openxmlformats-"
                "officedocument.spreadsheetml.sheet"
            ),
            use_container_width=True
        )

    else:

        st.info(
            "No generated Excel is currently available."
        )


# ============================================================
#                    FOOTER
# ============================================================

st.divider()

st.caption(
    "6thSense (6S-FO200) Vardaan"
)
